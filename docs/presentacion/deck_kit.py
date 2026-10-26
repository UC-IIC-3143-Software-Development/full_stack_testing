"""Kit de dibujo para el deck: estilos del curso y helpers sobre python-pptx.

Los estilos salen de las presentaciones del profesor (04-Testing y 06-Testing Frontend):
- 16:9 de 10 x 5,625 pulgadas, fondo blanco.
- Títulos en Bitter negrita (vienen del placeholder del master) y texto en Inter.
- Separadores de sección con texto azul #4A86E8 centrado.
- Diagramas con cajas redondeadas de colores planos y texto blanco en negrita.
- Código en Consolas con colores tipo VS Code claro.
"""

from __future__ import annotations

import copy
import math
import re
from dataclasses import dataclass, field

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# Paleta del curso (tema de 04-Testing más los colores que el profesor usa en diagramas)
INK = "1F2328"
BLACK = "000000"
GRAY = "57606A"
MUTED = "8C959F"
TEAL = "003D4F"  # accent1
PINK = "F2617A"  # accent2 (viñetas)
ORANGE = "CC850A"  # accent3
SAGE = "6B9E78"  # accent4
AQUA = "47A1AD"  # accent5
PLUM = "634F7D"  # accent6
BLUE = "4A86E8"  # separadores de sección
PURPLE = "9900FF"
DPURPLE = "351C75"
GREEN = "38761D"
RED = "980000"
LIGHT = "EDF1F3"  # lt2
LIGHTBLUE = "DCE9F7"
CODEBG = "F6F8FA"
BORDER = "D0D7DE"
WHITE = "FFFFFF"

FONT_BODY = "Inter"
FONT_TITLE = "Bitter"
FONT_CODE = "Consolas"

SLIDE_W = 10.0
SLIDE_H = 5.625

TAGS = {"r": RED, "g": GREEN, "b": BLUE, "t": TEAL, "o": ORANGE, "p": PURPLE, "k": PINK, "m": GRAY}
TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[a-z]\].+?\[/\])")

# Tamaños mínimos para proyectar (los verifica build_deck.check_sizes)
MIN_BODY = 12  # texto de cuerpo, tablas y etiquetas de diagramas
MIN_CAPTION = 10  # fuentes, pies de imagen y notas al margen (formas con nombre "caption")
MIN_CODE = 10  # bloques de código

# Advertencias de texto que probablemente no cabe (se imprimen al construir)
WARNINGS: list[str] = []


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def parse_runs(text: str, base: dict) -> list[tuple[str, dict]]:
    """Mini markup: **negrita**, `código`, [r]rojo[/], [g]verde[/], [b]azul[/] ..."""
    out = []
    for part in TOKEN.split(text):
        if not part:
            continue
        st = dict(base)
        if part.startswith("**") and part.endswith("**"):
            st["bold"] = True
            part = part[2:-2]
        elif part.startswith("`") and part.endswith("`"):
            st["font"] = FONT_CODE
            st["color"] = st.get("code_color", "A31515")
            part = part[1:-1]
        elif re.match(r"\[[a-z]\]", part):
            st["bold"] = True
            st["color"] = TAGS[part[1]]
            part = part[3:-3]
        out.append((part, st))
    return out


def est_lines(text: str, width_in: float, size: float, font: str = FONT_BODY, bold: bool = False) -> int:
    factor = {FONT_BODY: 0.52, FONT_TITLE: 0.46, FONT_CODE: 0.60}.get(font, 0.55) + (0.03 if bold else 0)
    clean = re.sub(r"\*\*|`|\[[a-z]\]|\[/\]", "", text)
    per_line = max(1, int(width_in / (size * factor / 72)))
    total = 0
    for para in clean.split("\n"):
        words = para.split(" ")
        lines, cur = 1, 0
        for w in words:
            add = len(w) + (1 if cur else 0)
            if cur + add > per_line and cur:
                lines += 1
                cur = len(w)
            else:
                cur += add
        total += lines
    return total


def _strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _set_run(run, text, size, color, bold=False, italic=False, font=FONT_BODY):
    run.text = text
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", font)
    rpr.set("lang", "es-CL")


def _bullet_ppr(paragraph, level: int, color: str, size: float):
    pPr = paragraph._p.get_or_add_pPr()
    indent = 0.18 if level == 0 else 0.16
    marl = 0.2 + level * 0.22
    pPr.set("marL", str(int(Inches(marl))))
    pPr.set("indent", str(int(-Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buClr", "a:buSzPct", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(buClr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000" if level == 0 else "85000")
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "●" if level == 0 else "○")


def _no_bullet(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    if pPr.find(qn("a:buNone")) is None:
        etree.SubElement(pPr, qn("a:buNone"))


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def fill_text(
    tf,
    paras,
    size=12,
    color=INK,
    bold=False,
    align="l",
    font=FONT_BODY,
    bullets=False,
    space_after=4,
    line_spacing=None,
    bullet_color=PINK,
    italic=False,
):
    """Escribe párrafos en un text frame. Un párrafo que empieza con dos espacios es nivel 2."""
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for raw in paras:
        level = 1 if raw.startswith("  ") else 0
        text = raw.strip() if level else raw
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = ALIGN[align]
        if space_after is not None:
            p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        if bullets:
            _bullet_ppr(p, level, bullet_color, size)
        else:
            _no_bullet(p)
        sz = size - (1 if level else 0)
        for t, st in parse_runs(text, {"bold": bold, "color": color, "font": font}):
            r = p.add_run()
            _set_run(r, t, sz, st["color"], st.get("bold", False), italic, st.get("font", font))


def replace_text(tf, paras: list[str], size: float | None = None):
    """Reemplaza el texto de un placeholder existente conservando el formato de su primer run."""
    ps = tf._txBody.findall(qn("a:p"))
    proto_p = copy.deepcopy(ps[0])
    runs = proto_p.findall(qn("a:r"))
    proto_r = copy.deepcopy(runs[0]) if runs else None
    for p in ps:
        tf._txBody.remove(p)
    for text in paras:
        p = copy.deepcopy(proto_p)
        for r in p.findall(qn("a:r")):
            p.remove(r)
        r = copy.deepcopy(proto_r) if proto_r is not None else etree.SubElement(p, qn("a:r"))
        r.find(qn("a:t")).text = text
        rpr = r.find(qn("a:rPr"))
        if rpr is None:
            rpr = etree.Element(qn("a:rPr"))
            r.insert(0, rpr)
        rpr.set("lang", "es-CL")
        if size:
            rpr.set("sz", str(int(size * 100)))
        end = p.find(qn("a:endParaRPr"))
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)
        tf._txBody.append(p)


class Deck:
    def __init__(self, base_path: str):
        self.prs = Presentation(base_path)
        layouts = {l.name: l for l in self.prs.slide_masters[1].slide_layouts}
        self.L_TITLE_ONLY = layouts["TITLE_ONLY"]
        self.L_CENTER = layouts["TITLE_2"]
        self.sections: list[tuple[str, list[int]]] = []
        self.notes_count = 0

    # ---------- diapositivas ----------
    def section(self, name: str):
        self.sections.append((name, []))

    def _register(self, slide):
        if self.sections:
            self.sections[-1][1].append(slide.slide_id)

    def _slide_number(self, slide):
        for ph in slide.slide_layout.placeholders:
            if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
                el = copy.deepcopy(ph._element)
                ids = [int(e.get("id")) for e in slide.shapes._spTree.iter() if e.tag == qn("p:cNvPr")]
                el.find(".//" + qn("p:cNvPr")).set("id", str(max(ids + [1]) + 1))
                slide.shapes._spTree.append(el)

    def content(self, title: str, size: float = 24):
        s = self.prs.slides.add_slide(self.L_TITLE_ONLY)
        self._register(s)
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Inches(0.4), Inches(0.32), Inches(9.2), Inches(0.62)
        tf = t.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = title
        r.font.size = Pt(size)
        r.font.bold = True
        if est_lines(title, 9.0, size, FONT_TITLE, True) > 1:
            WARNINGS.append(f"Título largo: {title!r}")
        self._slide_number(s)
        return s

    def center(self, paras, size=30, align="c", color=BLACK, font=FONT_TITLE, bold=True):
        """Diapositiva de frase central, como las del profesor (layout TITLE_2 con fondo blanco)."""
        s = self.prs.slides.add_slide(self.L_CENTER)
        self._register(s)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(WHITE)
        ph = [p for p in s.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.CENTER_TITLE][0]
        ph.left, ph.top, ph.width, ph.height = Inches(0.8), Inches(0.5), Inches(8.4), Inches(4.4)
        tf = ph.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        first = True
        for item in paras:
            if isinstance(item, tuple):
                text, opts = item
            else:
                text, opts = item, {}
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = ALIGN[opts.get("align", align)]
            p.space_after = Pt(opts.get("space_after", 14))
            for t, st in parse_runs(text, {"bold": opts.get("bold", bold), "color": opts.get("color", color)}):
                r = p.add_run()
                _set_run(r, t, opts.get("size", size), st["color"], st.get("bold", False),
                         opts.get("italic", False), opts.get("font", st.get("font", font)))
        self._slide_number(s)
        return s

    def divider(self, part: str, title: str, subtitle: str):
        return self.center(
            [
                (part.upper(), {"size": 12, "color": MUTED, "font": FONT_BODY, "space_after": 10}),
                (title, {"size": 36, "color": BLUE, "space_after": 16}),
                (subtitle, {"size": 15, "color": GRAY, "font": FONT_BODY, "bold": False}),
            ]
        )

    def notes(self, slide, text: str):
        text = re.sub(r"\s*\n\s*", " ", text.strip())
        slide.notes_slide.notes_text_frame.text = text
        self.notes_count += 1

    # ---------- formas ----------
    def text(self, slide, x, y, w, h, paras, size=12, color=INK, bold=False, align="l", anchor="t",
             font=FONT_BODY, bullets=False, space_after=4, margin=0.04, italic=False, line_spacing=None,
             bullet_color=PINK, name=None, check=True, caption=False):
        """Caja de texto. Con caption=True es una nota al margen (fuente, pie de imagen): puede ir a 10 pt."""
        if caption:
            name = "caption"
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Inches(margin))
        tf.vertical_anchor = ANCHOR[anchor]
        fill_text(tf, paras, size, color, bold, align, font, bullets, space_after, line_spacing, bullet_color, italic)
        if name:
            tb.name = name
        if check:
            self._check_fit(paras, w - 2 * margin - (0.2 if bullets else 0), h - 2 * margin, size, font, bold,
                            space_after, line_spacing)
        return tb

    def _check_fit(self, paras, w, h, size, font, bold, space_after, line_spacing):
        if isinstance(paras, str):
            paras = [paras]
        ls = line_spacing or 1.0
        need = 0.0
        for i, p in enumerate(paras):
            n = est_lines(p, w, size, font, bold)
            need += n * size * 1.2 * ls / 72 + ((space_after or 0) / 72 if i < len(paras) - 1 else 0)
        if need > h + 0.08:
            WARNINGS.append(f"Texto que puede no caber ({need:.2f} > {h:.2f} in): {str(paras)[:70]!r}")

    def box(self, slide, x, y, w, h, text="", fill=TEAL, color=WHITE, size=12, bold=None, line=None,
            line_w=1.0, shape="round", align="c", anchor="m", radius=0.08, font=FONT_BODY, margin=0.06,
            dash=False, paras=None, bullets=False, space_after=2, name=None, check=True):
        kind = {
            "round": MSO_SHAPE.ROUNDED_RECTANGLE,
            "rect": MSO_SHAPE.RECTANGLE,
            "oval": MSO_SHAPE.OVAL,
            "chevron": MSO_SHAPE.CHEVRON,
            "pentagon": MSO_SHAPE.PENTAGON,
            "trapezoid": MSO_SHAPE.TRAPEZOID,
            "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
            "diamond": MSO_SHAPE.DIAMOND,
            "callout": MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,
            "can": MSO_SHAPE.CAN,
            "doc": MSO_SHAPE.FLOWCHART_DOCUMENT,
            "flow": MSO_SHAPE.FLOWCHART_ALTERNATE_PROCESS,
        }[shape]
        sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        _strip_style(sp)
        if shape == "round":
            sp.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)))
        if fill:
            sp.fill.solid()
            sp.fill.fore_color.rgb = rgb(fill)
        else:
            sp.fill.background()
        if line:
            sp.line.color.rgb = rgb(line)
            sp.line.width = Pt(line_w)
            if dash:
                sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            sp.line.fill.background()
        tf = sp.text_frame
        tf.word_wrap = True
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Inches(margin))
        tf.vertical_anchor = ANCHOR[anchor]
        if bold is None:
            bold = paras is None  # con varios párrafos, solo lo marcado con ** va en negrita
        content = paras if paras is not None else (text if text else None)
        if content:
            fill_text(tf, content, size, color, bold, align, font, bullets, space_after)
            if check:
                self._check_fit(content, w - 2 * margin - (0.2 if bullets else 0), h - 2 * margin, size, font,
                                bold, space_after, None)
        if name:
            sp.name = name
        return sp

    def line(self, slide, x1, y1, x2, y2, color=GRAY, width=1.5, head=True, dash=False, tail=False):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        _strip_style(c)
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(width)
        ln = c.line._get_or_add_ln()
        if dash:
            prst = etree.SubElement(ln, qn("a:prstDash"))
            prst.set("val", "dash")
        if tail:
            etree.SubElement(ln, qn("a:headEnd")).attrib.update({"type": "triangle", "w": "med", "len": "med"})
        if head:
            etree.SubElement(ln, qn("a:tailEnd")).attrib.update({"type": "triangle", "w": "med", "len": "med"})
        return c

    def circle_num(self, slide, x, y, d, n, fill=TEAL, size=12, color=WHITE):
        return self.box(slide, x, y, d, d, str(n), fill=fill, color=color, size=size, shape="oval", margin=0)

    def chip(self, slide, x, y, w, h, text, fill=TEAL, size=12, color=WHITE):
        return self.box(slide, x, y, w, h, text, fill=fill, color=color, size=size, radius=0.06, margin=0.03)

    def card(self, slide, x, y, w, h, fill=LIGHT, line=None, radius=0.1):
        return self.box(slide, x, y, w, h, "", fill=fill, line=line, radius=radius)

    def image(self, slide, path, x, y, w=None, h=None, border=BORDER, alt=""):
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None,
                                       Inches(h) if h else None)
        if border:
            pic.line.color.rgb = rgb(border)
            pic.line.width = Pt(0.75)
        if alt:
            pic._element.nvPicPr.cNvPr.set("descr", alt)
        return pic

    # ---------- código ----------
    def code(self, slide, x, y, w, h, src: str, lang="ts", title=None, size=None, check=True):
        """Caja de código. Con size=None elige el tamaño más grande (11 a 10 pt) que cabe en la caja.

        El mínimo para proyectar es 10 pt: si el extracto no cabe a 10 pt se avisa y hay que acortarlo.
        """
        self.box(slide, x, y, w, h, "", fill=CODEBG, line=BORDER, line_w=0.75, radius=0.06)
        top = y + 0.08
        if title:
            self.text(slide, x + 0.12, y + 0.05, w - 0.24, 0.26, title, size=10, color=MUTED, font=FONT_CODE,
                      check=False, margin=0)
            top = y + 0.34
        if size is None:
            src_lines = src.strip("\n").split("\n")
            longest = max(len(l) for l in src_lines)
            size = MIN_CODE
            for cand in (11, 10.5, 10):
                fits_w = longest * cand * 0.60 / 72 <= w - 0.24
                fits_h = len(src_lines) * cand * 1.2 / 72 <= y + h - top - 0.06
                if fits_w and fits_h:
                    size = cand
                    break
        tb = slide.shapes.add_textbox(Inches(x + 0.12), Inches(top), Inches(w - 0.24), Inches(y + h - top - 0.06))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.auto_size = None
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Inches(0))
        lines = src.strip("\n").split("\n")
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(0)
            p.line_spacing = 1.0
            _no_bullet(p)
            toks = highlight(ln, lang) or [(" ", INK, False)]
            for t, col, b in toks:
                r = p.add_run()
                _set_run(r, t, size, col, b, font=FONT_CODE)
        if check:
            max_chars = max(len(l) for l in lines)
            avail_w = w - 0.24
            need_w = max_chars * size * 0.60 / 72
            need_h = len(lines) * size * 1.2 / 72
            avail_h = y + h - top - 0.06
            if need_w > avail_w + 0.02:
                WARNINGS.append(f"Código ancho ({max_chars} car., {need_w:.2f} > {avail_w:.2f} in): {lines[0][:40]!r}")
            if need_h > avail_h + 0.02:
                WARNINGS.append(f"Código alto ({len(lines)} líneas, {need_h:.2f} > {avail_h:.2f} in): {lines[0][:40]!r}")
        return tb

    # ---------- tablas ----------
    def table(self, slide, x, y, w, rows, col_w, row_h=0.36, size=12, header_fill=TEAL, header_color=WHITE,
              zebra=True, first_col_bold=False, col_colors=None, header_size=None, align=None):
        nrows, ncols = len(rows), len(rows[0])
        gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
        tbl = gf.table
        # sin estilo de tabla del tema: todo explícito
        tblPr = tbl._tbl.tblPr
        for attr in ("bandRow", "firstRow"):
            tblPr.set(attr, "0")
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(w * cw / total)
        for r in range(nrows):
            tbl.rows[r].height = Inches(row_h)
            for c in range(ncols):
                cell = tbl.cell(r, c)
                cell.margin_left = cell.margin_right = Inches(0.06)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                txt = rows[r][c]
                head = r == 0 and header_fill is not None
                fill = header_fill if head else (("F3F6F8" if r % 2 == 0 else WHITE) if zebra else WHITE)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(fill)
                col = header_color if head else (col_colors[c] if col_colors and col_colors[c] else INK)
                b = head or (first_col_bold and c == 0)
                tf = cell.text_frame
                tf.word_wrap = True
                fill_text(tf, txt.split("\n"), (header_size or size) if head else size, col, b,
                          (align[c] if align else "l"), space_after=0)
                _cell_borders(cell, BORDER)
        return gf


def _cell_borders(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    # Los bordes van antes del relleno dentro de tcPr
    fill = tcPr.find(qn("a:solidFill"))
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.Element(qn(tag), w=str(int(Pt(0.75))), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        if fill is not None:
            fill.addprevious(ln)
        else:
            tcPr.append(ln)


# ---------- resaltado de sintaxis simple ----------
KW = r"\b(?:const|let|var|async|new|class|interface|type|function|extends|implements|private|readonly|public|true|false|null|undefined|this|of|in|typeof|as)\b"
CTRL = r"\b(?:import|from|export|default|return|await|if|else|throw|for)\b"
FUNCS = r"\b(?:describe|it|test|expect|beforeAll|afterAll|beforeEach|afterEach|Given|When|Then|render|request)\b"
ID_IN_STR = re.compile(r"(\[(?:[A-Z]+-)+[A-Z0-9]+\])")

C_KW, C_CTRL, C_STR, C_COM, C_NUM, C_FN, C_KEY, C_TAG = (
    "0000FF", "AF00DB", "A31515", "008000", "098658", "795E26", "0451A5", "800080")


def _split_string(tok: str) -> list[tuple[str, str, bool]]:
    parts = []
    for piece in ID_IN_STR.split(tok):
        if not piece:
            continue
        if ID_IN_STR.fullmatch(piece):
            parts.append((piece, PURPLE, True))
        else:
            parts.append((piece, C_STR, False))
    return parts


def highlight(line: str, lang: str) -> list[tuple[str, str, bool]]:
    if lang in ("ts", "js", "tsx"):
        pat = re.compile(
            r"(?P<com>//.*$)|(?P<str>'(?:[^'\\]|\\.)*'|\"(?:[^\"\\]|\\.)*\"|`[^`]*`)|(?P<num>\b\d+(?:\.\d+)?\b)"
            rf"|(?P<ctrl>{CTRL})|(?P<kw>{KW})|(?P<fn>{FUNCS})"
        )
        colors = {"com": C_COM, "num": C_NUM, "ctrl": C_CTRL, "kw": C_KW, "fn": C_FN}
    elif lang == "yaml":
        pat = re.compile(
            r"(?P<com>#.*$)|(?P<key>^\s*-?\s*[\w$.'/{}\[\]-]+(?=:(\s|$)))|(?P<str>'(?:[^'\\]|\\.)*'|\"(?:[^\"\\]|\\.)*\")"
            r"|(?P<num>\b\d+\b)"
        )
        colors = {"com": C_COM, "key": C_KEY, "num": C_NUM}
    elif lang == "json":
        pat = re.compile(r"(?P<key>\"[^\"]+\"(?=\s*:))|(?P<str>\"(?:[^\"\\]|\\.)*\")|(?P<num>\b\d+\b)")
        colors = {"key": C_KEY, "num": C_NUM}
    elif lang == "gherkin":
        pat = re.compile(
            r"(?P<com>#.*$)|(?P<tag>@[\w-]+)|(?P<kwb>^\s*(?:Característica|Escenario|Antecedentes|Dado|Cuando|Entonces|Y)\b)"
            r"|(?P<str>\"[^\"]*\")"
        )
        colors = {"com": C_COM, "tag": C_TAG, "kwb": C_KW}
    elif lang == "md":
        pat = re.compile(r"(?P<head>^#+ .*$)|(?P<bold>\*\*[^*]+\*\*)|(?P<str>`[^`]+`)|(?P<com>^\|.*$)")
        colors = {"head": TEAL, "bold": INK, "com": GRAY}
    elif lang == "sh":
        pat = re.compile(
            r"(?P<prompt>^\$ )|(?P<ok>[✓✔].*$|^OK: .*$)"
            r"|(?P<bad>[✗✖].*$|^\s*error .*$|^x \d.*$|^FALLA: .*$|^(?:HUÉRFANO|ID NO EXISTE|PRESUPUESTO) .*$)"
            r"|(?P<com>#.*$)|(?P<str>'[^']*'|\"[^\"]*\")")
        colors = {"prompt": MUTED, "ok": GREEN, "bad": RED, "com": C_COM}
    else:
        return [(line, INK, False)]
    out: list[tuple[str, str, bool]] = []
    pos = 0
    for m in pat.finditer(line):
        if m.start() > pos:
            out.append((line[pos:m.start()], INK, False))
        kind = m.lastgroup
        tok = m.group(0)
        if kind == "str":
            out.extend(_split_string(tok))
        elif kind in ("kwb", "head", "bold", "ok", "bad"):
            out.append((tok, colors[kind], True))
        else:
            out.append((tok, colors[kind], False))
        pos = m.end()
    if pos < len(line):
        out.append((line[pos:], INK, False))
    return [t for t in out if t[0]]


def add_sections(prs, sections):
    """Agrega las secciones de PowerPoint (p14:sectionLst) para navegar el deck."""
    import uuid

    P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    pres = prs.part._element
    ext_lst = pres.find(qn("p:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(pres, qn("p:extLst"))
    ext = etree.SubElement(ext_lst, qn("p:ext"))
    ext.set("uri", "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    sl = etree.SubElement(ext, "{%s}sectionLst" % P14, nsmap={"p14": P14})
    for name, ids in sections:
        sec = etree.SubElement(sl, "{%s}section" % P14)
        sec.set("name", name)
        sec.set("id", "{%s}" % str(uuid.uuid5(uuid.NAMESPACE_URL, name)).upper())
        lst = etree.SubElement(sec, "{%s}sldIdLst" % P14)
        for i in ids:
            etree.SubElement(lst, "{%s}sldId" % P14).set("id", str(i))
